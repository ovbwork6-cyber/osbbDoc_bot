"""Resident, organization administration and Superadmin bot flows."""

from __future__ import annotations

import asyncio
import io
import logging
import re
from datetime import date, datetime, timedelta, timezone
from html import escape
from typing import Any
from zoneinfo import ZoneInfo

from aiogram import F, types
from aiogram.filters import Command
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import BufferedInputFile, InlineKeyboardButton, InlineKeyboardMarkup, KeyboardButton, ReplyKeyboardMarkup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import db_layer as db
from access import SUPERADMIN_ID, can_access_osbb, can_manage_osbb, can_process_osbb, is_superadmin, roles_for

logger = logging.getLogger("osbb_portal")
LOCAL_TZ = ZoneInfo("Europe/Kyiv")
METRICS = {
    "electricity": "Електроенергія",
    "cold_water": "Холодна вода",
    "hot_water": "Гаряча вода",
    "heating": "Опалення",
    "people": "Кількість осіб",
}


class PortalForm(StatesGroup):
    join_code = State()
    org_pick_apartment = State()
    apartment_pin = State()
    apartment_count = State()
    apartment_rename = State()
    apartment_info = State()
    reading_value = State()
    reading_photo = State()
    reading_review_note = State()
    resident_period = State()
    document_name = State()
    document_file = State()
    document_tags = State()
    apartment_doc_name = State()
    apartment_doc_file = State()
    resident_period_act = State()
    contact_value = State()
    role_grant = State()
    org_create = State()
    branch_create = State()
    payment_change = State()
    export_year = State()


def _kb(rows: list[list[str]], back: str = "p:home") -> InlineKeyboardMarkup:
    result = [[InlineKeyboardButton(text=label, callback_data=callback)] for row in rows for label, callback in [tuple(row)]]
    result.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=back)])
    return InlineKeyboardMarkup(inline_keyboard=result)


def _reply(rows: list[list[str]]) -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(keyboard=[[KeyboardButton(text=x) for x in row] for row in rows], resize_keyboard=True)


def _org(org_id: int) -> dict[str, Any] | None:
    return db.fetch_one("SELECT o.*,b.status branch_status,b.payment_status FROM organizations o JOIN branches b ON b.id=o.branch_id WHERE o.id=?", (int(org_id),))


def _org_by_name(name: str) -> dict[str, Any] | None:
    return db.fetch_one("SELECT * FROM organizations WHERE name=? AND active=1", (name,))


def _apartment_menu(apartment_id: int) -> InlineKeyboardMarkup:
    settings = db.fetch_all("SELECT metric,enabled FROM meter_settings WHERE apartment_id=?", (apartment_id,))
    orgstate=db.fetch_one("SELECT o.meters_enabled FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?",(apartment_id,))
    enabled = {r["metric"] for r in settings if r["enabled"]} if orgstate and orgstate["meters_enabled"] else set()
    rows = [[InlineKeyboardButton(text="📥 Подати показники", callback_data=f"p:submitlist:{apartment_id}")]] if enabled else []
    rows.extend([
        [InlineKeyboardButton(text="🗂 Архів показників", callback_data=f"p:readingarchive:{apartment_id}")],
        [InlineKeyboardButton(text="📚 Документи квартири", callback_data=f"p:residentdocs:{apartment_id}")],
        [InlineKeyboardButton(text="⬅️ До організацій", callback_data="p:home")],
    ])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def _cabinet_text(apartment_id: int) -> str:
    row = db.fetch_one("""SELECT a.number,a.resident_info,a.resident_count,o.name org FROM apartments a
        JOIN organizations o ON o.id=a.organization_id WHERE a.id=?""", (apartment_id,))
    if not row:
        return "Квартиру не знайдено."
    today = datetime.now(LOCAL_TZ).date().isoformat()
    active_period = db.fetch_one("SELECT person_count,starts_on,ends_on FROM resident_periods WHERE apartment_id=? AND ? BETWEEN starts_on AND ends_on ORDER BY starts_on DESC,id DESC LIMIT 1", (apartment_id, today))
    people = active_period["person_count"] if active_period else row["resident_count"]
    text = f"🏠 {row['org']} · квартира {row['number']}"
    if row["resident_info"]:
        text += f"\n\nℹ️ {row['resident_info']}"
    if people is not None:
        text += f"\n\n👥 Кількість осіб: {people}"
    if active_period:
        text += f" (з {active_period['starts_on']} до {active_period['ends_on']})"
    return text


def _can_submit_reading(org_id: int) -> bool:
    row = db.fetch_one("SELECT o.active,b.status branch_status FROM organizations o JOIN branches b ON b.id=o.branch_id WHERE o.id=?", (org_id,))
    return bool(row and row["active"] and row["branch_status"] == "active")


def _build_readings_xlsx(org_id: int, year: int) -> bytes:
    org = _org(org_id)
    if not org:
        raise ValueError("Organization does not exist")
    workbook = Workbook()
    first = workbook.active
    workbook.remove(first)
    header = ["Квартира", "Ключ входу", "Кількість осіб"]
    for metric in ("electricity", "cold_water", "hot_water", "heating"):
        header += [f"{METRICS[metric]} · значення", f"{METRICS[metric]} · статус"]
    for month in range(1, 13):
        period = f"{year}-{month:02d}"
        ws = workbook.create_sheet(f"{month:02d}")
        ws.append([f"{org['name']} · {METRICS['electricity']}, вода, опалення та мешканці · {period}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
        ws["A1"].font = Font(bold=True, color="FFFFFF", size=13)
        ws["A1"].fill = PatternFill("solid", fgColor="145C73")
        ws["A1"].alignment = Alignment(vertical="center")
        ws.append(header)
        for cell in ws[2]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="247B91")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        apartments = db.fetch_all("SELECT id,number,resident_pin,resident_count FROM apartments WHERE organization_id=? AND active=1 ORDER BY number", (org_id,))
        for apt in apartments:
            period_count = db.fetch_one("SELECT person_count FROM resident_periods WHERE apartment_id=? AND starts_on<=? AND ends_on>=? ORDER BY starts_on DESC,id DESC LIMIT 1", (apt["id"], f"{period}-31", f"{period}-01"))
            values: list[Any] = [apt["number"], apt["resident_pin"] or "", period_count["person_count"] if period_count else (apt["resident_count"] if apt["resident_count"] is not None else "")]
            for metric in ("electricity", "cold_water", "hot_water", "heating"):
                row = db.fetch_one("SELECT value,status FROM readings WHERE apartment_id=? AND metric=? AND period=?", (apt["id"], metric, period))
                values.extend([row["value"] if row else "", {"pending":"Очікує","accepted":"Прийнято","rejected":"Відхилено"}.get(row["status"], "") if row else ""])
            ws.append(values)
        ws.freeze_panes = "A3"
        ws.auto_filter.ref = f"A2:{chr(64+len(header))}{max(2,ws.max_row)}"
        ws.row_dimensions[1].height = 24
        ws.row_dimensions[2].height = 34
        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 19
        for col in range(4, len(header)+1):
            ws.column_dimensions[chr(64+col)].width = 21
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _build_pin_xlsx(org_id: int) -> bytes:
    org = _org(org_id)
    if not org:
        raise ValueError("Organization does not exist")
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Ключі квартир"
    ws.append([org["name"], "Квартира", "Ключ входу", "Прив’язаних Telegram акаунтів", "Кількість осіб"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="145C73")
    rows = db.fetch_all("""SELECT a.id,a.number,a.resident_pin,a.resident_count,
        (SELECT count(*) FROM apartment_members am WHERE am.apartment_id=a.id) members
        FROM apartments a WHERE a.organization_id=? AND a.active=1 ORDER BY a.number""", (org_id,))
    for row in rows:
        ws.append([org["name"], row["number"], row["resident_pin"] or "", row["members"], row["resident_count"] if row["resident_count"] is not None else ""])
    for col, width in {"A": 24, "B": 16, "C": 18, "D": 28, "E": 20}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{max(ws.max_row,1)}"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


async def resident_period_reminder_worker(bot: Any) -> None:
    """Send each start/end reminder once per recipient, using Kyiv dates."""
    processed_day: date | None = None
    while True:
        now = datetime.now(LOCAL_TZ)
        today = now.date()
        if processed_day != today:
            start_s = today.isoformat()
            end_s = (today + timedelta(days=7)).isoformat()
            periods = db.fetch_all("""SELECT p.id,p.apartment_id,p.person_count,p.starts_on,p.ends_on,a.number,o.name org,o.id org_id
                FROM resident_periods p JOIN apartments a ON a.id=p.apartment_id JOIN organizations o ON o.id=a.organization_id
                JOIN branches b ON b.id=o.branch_id WHERE o.active=1 AND b.status='active'
                AND (p.starts_on=? OR p.ends_on=?)""", (start_s, end_s))
            for period in periods:
                kind = "start" if period["starts_on"] == start_s else "end7"
                if kind == "start":
                    message = f"З {period['starts_on']} у квартирі {period['number']} ({period['org']}) діє кількість осіб: {period['person_count']}. Перевірте кабінет і документ-підставу."
                else:
                    message = f"Через 7 днів завершується період кількості осіб у квартирі {period['number']} ({period['org']}). Звірте актуальність даних."
                recipients = {int(r["telegram_id"]) for r in db.fetch_all("SELECT telegram_id FROM apartment_members WHERE apartment_id=?", (period["apartment_id"],))}
                admins = db.fetch_all("SELECT telegram_id FROM memberships WHERE organization_id=? AND active=1 AND role IN ('admin','fop')", (period["org_id"],))
                recipients.update(int(r["telegram_id"]) for r in admins)
                for recipient in recipients:
                    prior = db.fetch_one("SELECT 1 FROM period_reminder_logs WHERE period_id=? AND reminder_kind=? AND telegram_id=?", (period["id"], kind, recipient))
                    if prior:
                        continue
                    try:
                        await bot.send_message(recipient, message)
                        db.execute("INSERT OR IGNORE INTO period_reminder_logs(period_id,reminder_kind,telegram_id,sent_at) VALUES(?,?,?,?)", (period["id"], kind, recipient, db.utc_now()))
                    except Exception:
                        logger.exception("Unable to send resident period reminder to %s", recipient)
            processed_day = today
        tomorrow = datetime.combine(today + timedelta(days=1), datetime.min.time(), LOCAL_TZ)
        await asyncio.sleep(max(10, (tomorrow - datetime.now(LOCAL_TZ)).total_seconds()))


async def resident_documents_route(cb: types.CallbackQuery, org_id: int, bot: Any, apartment_id: int | None = None) -> None:
    org = _org(org_id)
    if not org or not _can_read_portal(cb.from_user.id, org["name"]):
        await cb.answer("Немає доступу.", show_alert=True)
        return
    if apartment_id is None:
        member = db.fetch_one("SELECT apartment_id FROM apartment_members WHERE telegram_id=? AND apartment_id IN (SELECT id FROM apartments WHERE organization_id=?) ORDER BY apartment_id LIMIT 1", (cb.from_user.id, org_id))
        apartment_id = int(member["apartment_id"]) if member else None
    rows = db.fetch_all(
        "SELECT id,name,file_id FROM portal_documents WHERE organization_id=? AND published=1 AND for_user=1 AND (apartment_id IS NULL OR apartment_id=?) ORDER BY created_at DESC",
        (org_id, apartment_id if apartment_id else -1),
    )
    if not rows:
        await cb.message.answer("Опублікованих документів для мешканців поки немає.")
    else:
        for row in rows:
            try:
                await bot.send_document(cb.message.chat.id, row["file_id"], caption=row["name"])
            except Exception:
                logger.exception("Could not send portal document id=%s", row["id"])
    await cb.answer()


def _audit(user_id: int, action: str, kind: str, entity_id: int | None = None, org: str | None = None, details: str = "") -> None:
    db.execute("INSERT INTO audit_log(user_id,action,entity_type,entity_id,osbb,details,created_at) VALUES(?,?,?,?,?,?,?)",
               (int(user_id), action, kind, entity_id, org, details, db.utc_now()))


def _is_manager(user_id: int, org: str) -> bool:
    return can_manage_osbb(user_id, org)


def _is_admin(user_id: int, org: str) -> bool:
    return can_access_osbb(user_id, org) and (is_superadmin(user_id) or "admin" in roles_for(user_id, org)) and can_manage_osbb(user_id, org)


def _can_export(user_id: int, org: str) -> bool:
    return can_access_osbb(user_id, org) and (is_superadmin(user_id) or bool(roles_for(user_id, org) & {"admin", "fop", "accountant"}))


def _can_export_keys(user_id: int, org: str) -> bool:
    return can_access_osbb(user_id, org) and (is_superadmin(user_id) or bool(roles_for(user_id, org) & {"admin", "accountant"}))


def _can_export_indicators(user_id: int, org: str) -> bool:
    return can_access_osbb(user_id, org) and (is_superadmin(user_id) or bool(roles_for(user_id, org) & {"admin", "accountant"}))


def _can_read_portal(user_id: int, org: str) -> bool:
    return can_access_osbb(user_id, org)


def _utcnow() -> datetime:
    return datetime.now(timezone.utc)


def _rate_limited(user_id: int, code_kind: str) -> tuple[bool, int]:
    row = db.fetch_one("SELECT failures,blocked_until FROM code_attempts_osbb WHERE telegram_id=? AND code_kind=?", (user_id, code_kind))
    if not row or not row["blocked_until"]:
        return False, 0
    try:
        until = datetime.fromisoformat(row["blocked_until"])
    except ValueError:
        return False, 0
    seconds = int((until - _utcnow()).total_seconds())
    if seconds <= 0:
        db.execute("UPDATE code_attempts_osbb SET failures=0,blocked_until=NULL WHERE telegram_id=? AND code_kind=?", (user_id, code_kind))
    return seconds > 0, max(seconds, 0)


def _failed_code(user_id: int, code_kind: str) -> int:
    row = db.fetch_one("SELECT failures FROM code_attempts_osbb WHERE telegram_id=? AND code_kind=?", (user_id, code_kind))
    failures = int(row["failures"] or 0) + 1 if row else 1
    blocked = (_utcnow() + timedelta(minutes=5)).isoformat(timespec="seconds") if failures >= 5 else None
    db.execute("""INSERT INTO code_attempts_osbb(telegram_id,code_kind,failures,blocked_until) VALUES(?,?,?,?)
        ON CONFLICT(telegram_id,code_kind) DO UPDATE SET failures=excluded.failures,blocked_until=excluded.blocked_until""",
        (user_id, code_kind, failures, blocked))
    return failures


def _clear_failures(user_id: int, code_kind: str) -> None:
    db.execute("DELETE FROM code_attempts_osbb WHERE telegram_id=? AND code_kind=?", (user_id, code_kind))


def _upsert_user(user: types.User) -> None:
    name = " ".join(x for x in (user.first_name, user.last_name) if x).strip()
    db.execute("""INSERT INTO users(telegram_id,full_name,username) VALUES(?,?,?)
        ON CONFLICT(telegram_id) DO UPDATE SET full_name=excluded.full_name,username=excluded.username,updated_at=CURRENT_TIMESTAMP""",
        (user.id, name, user.username or ""))


def _user_roles(user_id: int) -> set[str]:
    return roles_for(user_id)


def _main_menu(user_id: int) -> ReplyKeyboardMarkup:
    roles = _user_roles(user_id)
    if is_superadmin(user_id):
        return _reply([["📊 Dashboard", "📄 Акти"], ["🧾 Чеки ОСББ", "📈 Звіт по ОСББ"], ["🛠️ План робіт", "💰 Зарплати"], ["📁 Документи", "🏘️ Організації"], ["👥 Мешканці та будинок", "📥 Показники"], ["🔑 Ключі квартир", "📇 Контакти"], ["📝 Шаблони документів", "ℹ️ Інфо", "/admin"]])
    if roles & {"admin", "fop"}:
        rows=[["📊 Dashboard", "📄 Акти"], ["🧾 Чеки ОСББ", "📈 Звіт по ОСББ"], ["🛠️ План робіт", "💰 Зарплати"], ["📁 Документи", "🏘️ Мешканці та будинок"], ["📥 Показники", "🔑 Ключі квартир"], ["📇 Контакти", "📝 Шаблони документів"], ["ℹ️ Інфо"]]
        if "viewer" in roles: rows.append(["🔄 Оновити статус"])
        if roles & {"user","leadership"}: rows.append(["🏘️ Мої ОСББ","🔐 Ввести код ОСББ"] if "user" in roles else ["🏘️ Мої ОСББ"])
        return _reply(rows)
    if "accountant" in roles:
        return _reply([["📊 Dashboard", "📄 Акти"], ["🧾 Чеки ОСББ", "📈 Звіт по ОСББ"], ["📁 Документи", "📥 Показники"], ["🔑 Ключі квартир", "💰 Зарплати"], ["🛠️ План робіт", "ℹ️ Інфо"]])
    if "viewer" in roles:
        return _reply([["🔄 Оновити статус"]])
    if roles & {"leadership", "user"}:
        rows=[["🏘️ Мої ОСББ"]]
        if "user" in roles: rows.append(["🔐 Ввести код ОСББ"])
        rows += [["📚 Каталог документів", "🛠️ План робіт користувача"], ["📇 Контакти", "📝 Шаблони документів"], ["ℹ️ Інфо"]]
        return _reply(rows)
    return _reply([["🔐 Ввести код ОСББ"], ["ℹ️ Інфо"]])


def main_menu_for_user(user_id: int) -> ReplyKeyboardMarkup:
    return _main_menu(user_id)


def upsert_user(user: types.User) -> None:
    _upsert_user(user)


def _org_keyboard(user_id: int, prefix: str, roles: set[str] | None = None) -> InlineKeyboardMarkup:
    if is_superadmin(user_id):
        rows = db.fetch_all("SELECT id,name FROM organizations WHERE active=1 ORDER BY name")
    else:
        sql = "SELECT DISTINCT o.id,o.name FROM organizations o JOIN memberships m ON m.organization_id=o.id JOIN branches b ON b.id=o.branch_id WHERE m.telegram_id=? AND m.active=1 AND o.active=1 AND b.status='active'"
        params: list[Any] = [user_id]
        if roles:
            placeholders = ",".join("?" for _ in roles)
            sql += f" AND m.role IN ({placeholders})"
            params.extend(sorted(roles))
        sql += " ORDER BY o.name"
        rows = db.fetch_all(sql, params)
    buttons = [[InlineKeyboardButton(text=r["name"], callback_data=f"p:{prefix}:{r['id']}")] for r in rows]
    buttons.append([InlineKeyboardButton(text="⬅️ Меню", callback_data="p:home")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def _is_manager_member(user_id: int, org_id: int) -> bool:
    org = _org(org_id)
    return bool(org and can_manage_osbb(user_id, org["name"]))


def _can_process_member(user_id: int, org_id: int) -> bool:
    org = _org(org_id)
    return bool(org and can_process_osbb(user_id, org["name"]))


async def _send_home(message: types.Message, user_id: int, text: str = "Оберіть потрібний розділ.") -> None:
    await message.answer(text, reply_markup=_main_menu(user_id))


def register_handlers(dp: Any, bot: Any) -> None:
    @dp.message(Command("admin"))
    async def superadmin_home(m: types.Message, state: FSMContext) -> None:
        if not is_superadmin(m.from_user.id):
            return await m.answer("⛔ Команда доступна лише Superadmin.")
        await state.clear()
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🌿 Гілки й оплата", callback_data="p:sa:branches")],
            [InlineKeyboardButton(text="🏘️ ОСББ та кооперативи", callback_data="p:sa:orgs")],
            [InlineKeyboardButton(text="👥 Користувачі та ролі", callback_data="p:sa:users")],
            [InlineKeyboardButton(text="🧾 Журнал аудиту", callback_data="p:sa:audit")],
            [InlineKeyboardButton(text="⬅️ Робоче меню", callback_data="p:home")],
        ])
        await m.answer("Керування системою", reply_markup=kb)

    @dp.callback_query(F.data == "p:home")
    async def go_home(cb: types.CallbackQuery, state: FSMContext) -> None:
        await state.clear()
        await cb.message.answer("Робоче меню", reply_markup=_main_menu(cb.from_user.id))
        await cb.answer()

    @dp.message(F.text == "ℹ️ Інфо")
    async def info(m: types.Message) -> None:
        await m.answer("ОСББ/кооперативи: підключайтеся кодом, який видає Admin. Якщо доступ втрачено, зверніться до голови організації. Підтримка: звернення через Superadmin.")

    @dp.message(F.text == "📝 Шаблони документів")
    async def templates_stub(m: types.Message) -> None:
        await m.answer("Модуль шаблонів документів готується.")

    @dp.message(F.text == "🔐 Ввести код ОСББ")
    async def join_code_start(m: types.Message, state: FSMContext) -> None:
        await state.clear()
        await state.set_state(PortalForm.join_code)
        await m.answer("Введіть код мешканця або код члена правління, який надав Admin.")

    @dp.message(PortalForm.join_code)
    async def join_code_submit(m: types.Message, state: FSMContext) -> None:
        user_id = m.from_user.id
        value = (m.text or "").strip().upper()
        blocked, seconds = _rate_limited(user_id, "join")
        if blocked:
            await state.clear()
            return await m.answer(f"Забагато помилок. Спробуйте знову через {max(1, (seconds + 59)//60)} хв.")
        matches = db.fetch_all("SELECT id,name,resident_code,leadership_code FROM organizations WHERE active=1")
        match = next((r for r in matches if value and len(value)<=40 and value in {str(r["resident_code"] or "").upper(), str(r["leadership_code"] or "").upper()}), None)
        kind = "leadership" if match and value == str(match["leadership_code"] or "").upper() else "user"
        if not match:
            failures = _failed_code(user_id, "join")
            await state.clear()
            if failures >= 5:
                return await m.answer("П’ять спроб використано. Введення кодів заблоковано на 5 хвилин. Зверніться до голови ОСББ.")
            return await m.answer(f"Код не знайдено. Залишилося спроб: {5-failures}.")
        org = _org(int(match["id"]))
        if not org or org["branch_status"] != "active":
            await state.clear()
            return await m.answer("Організація тимчасово недоступна. Зверніться до Admin.")
        _upsert_user(m.from_user)
        db.execute("INSERT OR IGNORE INTO memberships(telegram_id,organization_id,role) VALUES(?,?,?)", (user_id, int(match["id"]), kind))
        _clear_failures(user_id, "join")
        _audit(user_id, "join_by_code", "organization", int(match["id"]), match["name"], f"role={kind}")
        await state.clear()
        await m.answer(f"Доступ до {match['name']} додано. Роль: {'Leadership' if kind=='leadership' else 'мешканець'}.", reply_markup=_main_menu(user_id))

    @dp.message(F.text == "🏘️ Мої ОСББ")
    async def my_organizations(m: types.Message) -> None:
        roles = _user_roles(m.from_user.id)
        if not roles & {"user", "leadership"}:
            return await m.answer("Підключених організацій поки немає. Введіть код, який надав Admin.", reply_markup=_main_menu(m.from_user.id))
        await m.answer("Оберіть організацію:", reply_markup=_org_keyboard(m.from_user.id, "resident", {"user", "leadership"}))

    @dp.callback_query(F.data.startswith("p:resident:"))
    async def resident_org(cb: types.CallbackQuery, state: FSMContext) -> None:
        org_id = int(cb.data.rsplit(":", 1)[1])
        org = _org(org_id)
        if not org or not _can_read_portal(cb.from_user.id, org["name"]):
            return await cb.answer("Немає доступу до організації.", show_alert=True)
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔑 Кабінет квартири", callback_data=f"p:cabinet:{org_id}")],
            [InlineKeyboardButton(text="📚 Каталог документів", callback_data=f"p:documents:{org_id}")],
            [InlineKeyboardButton(text="🛠️ План робіт", callback_data=f"p:jobs:{org_id}")],
            [InlineKeyboardButton(text="📇 Контакти", callback_data=f"p:contacts:{org_id}")],
        ])
        await cb.message.answer(f"{org['name']} — оберіть розділ.", reply_markup=kb)
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:cabinet:"))
    async def apartment_choose(cb: types.CallbackQuery, state: FSMContext) -> None:
        org_id = int(cb.data.rsplit(":", 1)[1])
        org = _org(org_id)
        if not org or not _can_read_portal(cb.from_user.id, org["name"]):
            return await cb.answer("Немає доступу.", show_alert=True)
        rows = db.fetch_all("SELECT id,number FROM apartments WHERE organization_id=? AND active=1 ORDER BY number", (org_id,))
        if not rows:
            return await cb.answer("Admin ще не додав квартири.", show_alert=True)
        buttons = [[InlineKeyboardButton(text=f"Квартира {r['number']}", callback_data=f"p:apt:{r['id']}")] for r in rows]
        await cb.message.answer(f"{org['name']}: оберіть свою квартиру.", reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
        await state.set_state(PortalForm.org_pick_apartment)
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:apt:"), PortalForm.org_pick_apartment)
    async def apartment_pin_prompt(cb: types.CallbackQuery, state: FSMContext) -> None:
        apartment_id = int(cb.data.rsplit(":", 1)[1])
        apt = db.fetch_one("SELECT id,number,organization_id,resident_pin,resident_pin_blocked_until FROM apartments WHERE id=? AND active=1", (apartment_id,))
        if not apt:
            await state.clear()
            return await cb.answer("Квартиру не знайдено.", show_alert=True)
        db.execute("INSERT OR IGNORE INTO users(telegram_id,full_name) VALUES(?,?)", (cb.from_user.id, cb.from_user.full_name))
        await state.update_data(apartment_id=apartment_id)
        await state.set_state(PortalForm.apartment_pin)
        prompt = "Створіть ключ квартири: введіть 4 цифри." if not apt["resident_pin"] else "Введіть 4-значний ключ квартири."
        await cb.message.answer(f"Квартира {apt['number']}. {prompt}")
        await cb.answer()

    @dp.message(PortalForm.apartment_pin)
    async def apartment_pin_submit(m: types.Message, state: FSMContext) -> None:
        data = await state.get_data()
        apt_id = int(data.get("apartment_id") or 0)
        apt = db.fetch_one("SELECT * FROM apartments WHERE id=? AND active=1", (apt_id,))
        if not apt:
            await state.clear()
            return await m.answer("Квартиру не знайдено. Почніть знову через меню.")
        now = _utcnow()
        if apt["resident_pin_blocked_until"]:
            try:
                remaining = int((datetime.fromisoformat(apt["resident_pin_blocked_until"]) - now).total_seconds())
            except ValueError:
                remaining = 0
            if remaining > 0:
                await state.clear()
                return await m.answer(f"Введення ключа заблоковано на {max(1,(remaining+59)//60)} хв. Зверніться до голови ОСББ, щоб уточнити ключ.")
            db.execute("UPDATE apartments SET resident_pin_failures=0,resident_pin_blocked_until=NULL WHERE id=?",(apt_id,))
            apt["resident_pin_failures"]=0
        pin = (m.text or "").strip()
        if not re.fullmatch(r"\d{4}", pin):
            failures = int(apt["resident_pin_failures"] or 0) + 1
            blocked_until = (now + timedelta(minutes=5)).isoformat(timespec="seconds") if failures >= 5 else None
            db.execute("UPDATE apartments SET resident_pin_failures=?,resident_pin_blocked_until=? WHERE id=?", (failures, blocked_until, apt_id))
            await state.clear()
            if failures >= 5:
                return await m.answer("П’ять спроб використано. Ключ квартири заблоковано на 5 хвилин. Зверніться до голови ОСББ, щоб уточнити ключ.")
            return await m.answer(f"Ключ має містити 4 цифри. Залишилося спроб: {5-failures}.")
        if not apt["resident_pin"]:
            db.execute("UPDATE apartments SET resident_pin=?,resident_pin_failures=0,resident_pin_blocked_until=NULL WHERE id=?", (pin, apt_id))
            _audit(m.from_user.id, "create_apartment_key", "apartment", apt_id)
        elif pin != apt["resident_pin"]:
            failures = int(apt["resident_pin_failures"] or 0) + 1
            blocked_until = (now + timedelta(minutes=5)).isoformat(timespec="seconds") if failures >= 5 else None
            db.execute("UPDATE apartments SET resident_pin_failures=?,resident_pin_blocked_until=? WHERE id=?", (failures, blocked_until, apt_id))
            await state.clear()
            if failures >= 5:
                return await m.answer("П’ять спроб використано. Ключ квартири заблоковано на 5 хвилин. Зверніться до голови ОСББ, щоб уточнити ключ.")
            return await m.answer(f"Ключ не підійшов. Залишилося спроб: {5-failures}.")
        else:
            db.execute("UPDATE apartments SET resident_pin_failures=0,resident_pin_blocked_until=NULL WHERE id=?", (apt_id,))
        db.execute("INSERT OR IGNORE INTO apartment_members(apartment_id,telegram_id) VALUES(?,?)", (apt_id, m.from_user.id))
        db.execute("UPDATE users SET updated_at=CURRENT_TIMESTAMP WHERE telegram_id=?", (m.from_user.id,))
        await state.clear()
        await m.answer(await _cabinet_text(apt_id), reply_markup=_apartment_menu(apt_id))

    @dp.callback_query(F.data.startswith("p:read:"))
    async def read_metric_start(cb: types.CallbackQuery, state: FSMContext) -> None:
        _, _, apt_raw, metric = cb.data.split(":", 3)
        apt_id = int(apt_raw)
        if not db.fetch_one("SELECT 1 FROM apartment_members WHERE apartment_id=? AND telegram_id=?", (apt_id, cb.from_user.id)) or not db.fetch_one("SELECT 1 FROM apartments WHERE id=? AND active=1",(apt_id,)):
            return await cb.answer("Спочатку введіть ключ квартири.", show_alert=True)
        settings = db.fetch_one("SELECT enabled,mode FROM meter_settings WHERE apartment_id=? AND metric=?", (apt_id, metric))
        orgstate=db.fetch_one("SELECT o.meters_enabled FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?",(apt_id,))
        if not orgstate or not orgstate["meters_enabled"]:
            return await cb.answer("Admin призупинив подачу показників.",show_alert=True)
        if metric == "people" or not settings or not settings["enabled"] or settings["mode"] == "people":
            return await cb.answer("Для цього показника подача зараз вимкнена.", show_alert=True)
        today = datetime.now(LOCAL_TZ).date()
        await state.update_data(apartment_id=apt_id, metric=metric, mode=settings["mode"], period=today.strftime("%Y-%m"))
        await state.set_state(PortalForm.reading_value)
        await cb.message.answer(f"Введіть показник «{METRICS[metric]}» за {today.strftime('%m.%Y')}.")
        await cb.answer()

    @dp.message(PortalForm.reading_value)
    async def reading_value_submit(m: types.Message, state: FSMContext) -> None:
        data = await state.get_data()
        raw = (m.text or "").strip().replace(",", ".")
        try:
            value = float(raw)
            if value < 0 or value > 10**12:
                raise ValueError
        except ValueError:
            return await m.answer("Введіть невід’ємне числове значення.")
        await state.update_data(value=str(raw))
        await state.set_state(PortalForm.reading_photo)
        await m.answer("Тепер надішліть фото показника лічильника для перевірки Admin.")

    @dp.message(PortalForm.reading_photo, F.photo)
    async def reading_photo_submit(m: types.Message, state: FSMContext) -> None:
        data = await state.get_data()
        apt_id = int(data["apartment_id"])
        apt = db.fetch_one("SELECT a.number,a.active,o.name org,o.id org_id FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?", (apt_id,))
        if not apt or not apt["active"] or not db.fetch_one("SELECT 1 FROM apartment_members WHERE apartment_id=? AND telegram_id=?", (apt_id, m.from_user.id)):
            await state.clear()
            return await m.answer("Квартиру не знайдено або доступ втрачено.")
        settings = db.fetch_one("SELECT enabled,mode FROM meter_settings WHERE apartment_id=? AND metric=?", (apt_id, data["metric"]))
        org_state = db.fetch_one("SELECT meters_enabled FROM organizations WHERE id=?", (apt["org_id"],))
        if not _can_submit_reading(apt["org_id"]) or not org_state or not org_state["meters_enabled"] or not settings or not settings["enabled"] or settings["mode"] == "people":
            await state.clear()
            return await m.answer("Подача показників наразі недоступна.")
        period = data["period"]
        existing=db.fetch_one("SELECT status FROM readings WHERE apartment_id=? AND metric=? AND period=?",(apt_id,data["metric"],period))
        if existing and existing["status"]=="accepted":
            await state.clear()
            return await m.answer("Показник за цей місяць уже прийнятий Admin. Якщо треба виправити його, зверніться до голови ОСББ.")
        db.execute("""INSERT INTO readings(apartment_id,metric,period,value,photo_file_id,status,submitted_by,created_at)
            VALUES(?,?,?,?,?,'pending',?,?) ON CONFLICT(apartment_id,metric,period) DO UPDATE SET
            value=excluded.value,photo_file_id=excluded.photo_file_id,status='pending',decision_note='',submitted_by=excluded.submitted_by,
            created_at=excluded.created_at,decided_by=NULL,decided_at=NULL""",
            (apt_id, data["metric"], period, data["value"], m.photo[-1].file_id, m.from_user.id, db.utc_now()))
        _audit(m.from_user.id, "submit_reading", "reading", None, apt["org"], f"apt={apt['number']}; metric={data['metric']}; period={period}")
        await state.clear()
        await m.answer("Показник надіслано Admin на перевірку. Фото буде прибрано з обліку після рішення.")

    @dp.message(PortalForm.reading_photo)
    async def reading_photo_wrong(m: types.Message) -> None:
        await m.answer("Для перевірки показника потрібно надіслати фото лічильника.")

    @dp.callback_query(F.data.startswith("p:resident_docs:"))
    async def resident_documents(cb: types.CallbackQuery) -> None:
        org_id = int(cb.data.rsplit(":", 1)[1])
        org = _org(org_id)
        if not org or not _can_read_portal(cb.from_user.id, org["name"]):
            return await cb.answer("Немає доступу.", show_alert=True)
        apartment = db.fetch_one("SELECT apartment_id FROM apartment_members WHERE telegram_id=? ORDER BY apartment_id LIMIT 1", (cb.from_user.id,))
        rows = db.fetch_all("SELECT id,name,file_id FROM portal_documents WHERE organization_id=? AND published=1 AND for_user=1 AND (apartment_id IS NULL OR apartment_id=?) ORDER BY created_at DESC", (org_id, int(apartment["apartment_id"]) if apartment else -1))
        if not rows:
            await cb.message.answer("Опублікованих документів для мешканців поки немає.")
        else:
            for row in rows:
                await bot.send_document(cb.message.chat.id, row["file_id"], caption=row["name"])
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:documents:"))
    async def documents_button(cb: types.CallbackQuery) -> None:
        org_id = int(cb.data.rsplit(":", 1)[1])
        await resident_documents_route(cb, org_id, bot)

    @dp.callback_query(F.data.startswith("p:residentdocs:"))
    async def apartment_documents(cb: types.CallbackQuery) -> None:
        apartment_id=int(cb.data.rsplit(":",1)[1])
        member = db.fetch_one("SELECT a.organization_id,a.id apartment_id FROM apartment_members am JOIN apartments a ON a.id=am.apartment_id WHERE am.telegram_id=? AND a.id=?", (cb.from_user.id,apartment_id))
        if not member:
            return await cb.answer("Спочатку відкрийте кабінет квартири.", show_alert=True)
        await resident_documents_route(cb, int(member["organization_id"]), bot, int(member["apartment_id"]))

    @dp.callback_query(F.data.startswith("p:submitlist:"))
    async def submit_list(cb: types.CallbackQuery) -> None:
        apt_id = int(cb.data.rsplit(":", 1)[1])
        if not db.fetch_one("SELECT 1 FROM apartment_members WHERE apartment_id=? AND telegram_id=?", (apt_id, cb.from_user.id)):
            return await cb.answer("Немає доступу до квартири.", show_alert=True)
        orgstate=db.fetch_one("SELECT o.meters_enabled FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?",(apt_id,))
        if not orgstate or not orgstate["meters_enabled"]: return await cb.answer("Admin призупинив подачу показників.",show_alert=True)
        rows = db.fetch_all("SELECT metric FROM meter_settings WHERE apartment_id=? AND enabled=1 AND mode='meter' AND metric!='people' ORDER BY metric", (apt_id,))
        if not rows:
            return await cb.answer("Admin не призначив показники для подачі.", show_alert=True)
        kb = InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text=METRICS[r['metric']], callback_data=f"p:read:{apt_id}:{r['metric']}")] for r in rows])
        await cb.message.answer("Оберіть показник:", reply_markup=kb)
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:readingarchive:"))
    async def reading_archive(cb: types.CallbackQuery) -> None:
        apt_id = int(cb.data.rsplit(":", 1)[1])
        if not db.fetch_one("SELECT 1 FROM apartment_members WHERE apartment_id=? AND telegram_id=?", (apt_id, cb.from_user.id)):
            return await cb.answer("Немає доступу до квартири.", show_alert=True)
        apt = db.fetch_one("SELECT number FROM apartments WHERE id=?", (apt_id,))
        rows = db.fetch_all("SELECT metric,period,value,status,decision_note FROM readings WHERE apartment_id=? ORDER BY period DESC,metric", (apt_id,))
        if not rows:
            await cb.message.answer("Історії показників поки немає.")
        else:
            labels = {"pending":"Очікує перевірки","accepted":"Прийнято","rejected":"Відхилено"}
            for r in rows[:60]:
                text=f"Квартира {apt['number']} · {r['period']}\n{METRICS.get(r['metric'],r['metric'])}: {r['value']}\nСтатус: {labels.get(r['status'],r['status'])}"
                if r["decision_note"]: text+=f"\nПояснення: {r['decision_note']}"
                await cb.message.answer(text)
        await cb.answer()

    @dp.message(F.text == "📊 Dashboard")
    async def dashboard(m: types.Message) -> None:
        if is_superadmin(m.from_user.id):
            orgs = db.fetch_all("SELECT o.id,o.name,b.payment_status,b.status branch_status FROM organizations o JOIN branches b ON b.id=o.branch_id WHERE o.active=1 ORDER BY o.name")
        else:
            orgs = db.fetch_all("SELECT DISTINCT o.id,o.name,b.payment_status,b.status branch_status FROM organizations o JOIN branches b ON b.id=o.branch_id JOIN memberships mm ON mm.organization_id=o.id WHERE mm.telegram_id=? AND mm.active=1 AND o.active=1 AND b.status='active' ORDER BY o.name", (m.from_user.id,))
        if not orgs:
            return await m.answer("Немає доступних організацій.")
        lines = ["📊 Dashboard"]
        for org in orgs:
            roles=roles_for(m.from_user.id,org["name"])
            counts=db.fetch_one("SELECT (SELECT count(*) FROM acts WHERE osbb=? AND status NOT IN ('Завершено!','Роботу завершено')) acts,(SELECT count(*) FROM docs WHERE osbb=? AND status NOT IN ('Завершено!','Роботу завершено')) docs,(SELECT count(*) FROM jobs WHERE osbb=? AND status NOT IN ('Роботу закінчено','Роботу завершено')) jobs",(org["name"],org["name"],org["name"]))
            pending=db.fetch_one("SELECT count(*) n FROM readings r JOIN apartments a ON a.id=r.apartment_id WHERE a.organization_id=? AND r.status='pending'",(org["id"],))["n"]
            today=datetime.now(LOCAL_TZ).date()
            due=db.fetch_one("SELECT count(*) n FROM resident_periods p JOIN apartments a ON a.id=p.apartment_id WHERE a.organization_id=? AND p.ends_on>=? AND p.starts_on<=?",(org["id"],today.isoformat(),(today+timedelta(days=30)).isoformat()))["n"]
            lines.append(f"\n🏘 {org['name']} · {org['payment_status']}\nАкти: {counts['acts']} · чеки: {counts['docs']} · роботи: {counts['jobs']}")
            if roles & {"admin","fop","accountant"}:
                lines.append(f"Показники на перевірку: {pending}")
            if roles & {"admin","fop"} and due:
                lines.append(f"Наближається зміна кількості мешканців: {due} квартир")
        kb=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text=f"📥 Показники · {o['name']}",callback_data=f"p:pending:{o['id']}")] for o in orgs if roles_for(m.from_user.id,o['name']) & {"admin","fop"}])
        await m.answer("\n".join(lines),reply_markup=kb if kb.inline_keyboard else None)

    @dp.message(F.text == "📥 Показники")
    async def readings_export_menu(m: types.Message) -> None:
        await m.answer("Оберіть організацію для вивантаження показників:", reply_markup=_org_keyboard(m.from_user.id,"export",{"admin","accountant"}))

    @dp.message(F.text == "🔑 Ключі квартир")
    async def apartment_keys_menu(m: types.Message) -> None:
        await m.answer("Оберіть організацію для вивантаження ключів квартир:", reply_markup=_org_keyboard(m.from_user.id,"pinexport",{"admin","accountant"}))

    @dp.message(F.text == "📁 Документи")
    @dp.message(F.text == "📚 Каталог документів")
    async def documents_menu(m: types.Message) -> None:
        roles=_user_roles(m.from_user.id)
        required={"admin","fop"} if roles & {"admin","fop"} else {"accountant","leadership","user"}
        await m.answer("Оберіть організацію:",reply_markup=_org_keyboard(m.from_user.id,"docmenu",required))

    @dp.callback_query(F.data.startswith("p:docmenu:"))
    async def organization_documents(cb: types.CallbackQuery) -> None:
        org_id=int(cb.data.rsplit(":",1)[1]); org=_org(org_id)
        if not org or not _can_read_portal(cb.from_user.id,org["name"]): return await cb.answer("Немає доступу.",show_alert=True)
        roles=roles_for(cb.from_user.id,org["name"])
        if roles & {"admin","fop"}:
            rows=db.fetch_all("SELECT id,name,published,for_user,for_leadership FROM portal_documents WHERE organization_id=? ORDER BY created_at DESC LIMIT 30",(org_id,))
            kb=[[InlineKeyboardButton(text=f"📄 {r['name']} · {'опубліковано' if r['published'] else 'чернетка'}",callback_data=f"p:docview:{r['id']}")] for r in rows]
            kb.append([InlineKeyboardButton(text="➕ Додати документ",callback_data=f"p:docadd:{org_id}")])
        else:
            visibility = []
            if "accountant" in roles: visibility.append("published=1")
            if "leadership" in roles: visibility.append("published=1 AND for_leadership=1")
            if "user" in roles: visibility.append("published=1 AND for_user=1")
            if not visibility: return await cb.answer("Немає доступу до документів.",show_alert=True)
            rows=db.fetch_all("SELECT id,name,file_id FROM portal_documents WHERE organization_id=? AND ("+" OR ".join(f"({v})" for v in visibility)+") ORDER BY created_at DESC",(org_id,))
            kb=[]
            for row in rows:
                try: await bot.send_document(cb.message.chat.id,row["file_id"],caption=row["name"])
                except Exception: logger.exception("Could not send Leadership document %s",row["id"])
            await cb.answer("Документи надіслано."); return
        kb.append([InlineKeyboardButton(text="⬅️ Назад",callback_data="p:home")])
        await cb.message.answer(f"Документи {org['name']}:",reply_markup=InlineKeyboardMarkup(inline_keyboard=kb))
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:docadd:"))
    async def document_add_start(cb: types.CallbackQuery,state: FSMContext) -> None:
        org_id=int(cb.data.rsplit(":",1)[1]); org=_org(org_id)
        if not org or not _is_manager(cb.from_user.id,org["name"]): return await cb.answer("Немає доступу.",show_alert=True)
        await state.update_data(org_id=org_id); await state.set_state(PortalForm.document_name)
        await cb.message.answer("Введіть назву документа; далі надішліть PDF або фото."); await cb.answer()

    @dp.message(PortalForm.document_name)
    async def document_name_save(m: types.Message,state: FSMContext) -> None:
        name=(m.text or "").strip()
        if not name or len(name)>150: return await m.answer("Назва має містити до 150 символів.")
        await state.update_data(document_name=name); await state.set_state(PortalForm.document_file)
        await m.answer("Надішліть PDF-файл або фото документа.")

    @dp.message(PortalForm.document_file)
    async def document_file_save(m: types.Message,state: FSMContext) -> None:
        data=await state.get_data(); org_id=int(data.get("org_id") or 0); org=_org(org_id)
        if not org or not _is_manager(m.from_user.id,org["name"]): await state.clear(); return await m.answer("Немає доступу.")
        if m.document:
            if m.document.mime_type!="application/pdf" and not (m.document.file_name or "").lower().endswith(".pdf"):
                return await m.answer("Потрібен PDF або фото.")
            file_id=m.document.file_id
        elif m.photo: file_id=m.photo[-1].file_id
        else: return await m.answer("Надішліть PDF або фото.")
        doc_id=db.execute("INSERT INTO portal_documents(organization_id,name,file_id,created_by) VALUES(?,?,?,?)",(org_id,data["document_name"],file_id,m.from_user.id))
        _audit(m.from_user.id,"create_portal_document","document",doc_id,org["name"])
        await state.clear()
        kb=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="📣 Опублікувати",callback_data=f"p:doctoggle:{doc_id}:published")],
            [InlineKeyboardButton(text="👥 Для мешканців",callback_data=f"p:doctoggle:{doc_id}:for_user")],
            [InlineKeyboardButton(text="🏛 Для Leadership",callback_data=f"p:doctoggle:{doc_id}:for_leadership")],
            [InlineKeyboardButton(text="⬅️ Готово",callback_data="p:home")],
        ])
        await m.answer(f"Документ «{data['document_name']}» збережено як чернетку. Виберіть видимість.",reply_markup=kb)

    @dp.callback_query(F.data.startswith("p:doctoggle:"))
    async def document_visibility(cb: types.CallbackQuery) -> None:
        _,_,doc_raw,flag=cb.data.split(":",3); doc_id=int(doc_raw)
        if flag not in {"published","for_user","for_leadership"}: return await cb.answer("Некоректне поле.",show_alert=True)
        row=db.fetch_one("SELECT d.*,o.name org FROM portal_documents d JOIN organizations o ON o.id=d.organization_id WHERE d.id=?",(doc_id,))
        if not row or not _is_manager(cb.from_user.id,row["org"]): return await cb.answer("Немає доступу.",show_alert=True)
        new=0 if row[flag] else 1
        db.execute(f"UPDATE portal_documents SET {flag}=? WHERE id=?",(new,doc_id))
        _audit(cb.from_user.id,"change_document_visibility","document",doc_id,row["org"],f"{flag}={new}")
        await cb.answer("Видимість оновлено.")

    @dp.callback_query(F.data.startswith("p:docview:"))
    async def document_view(cb: types.CallbackQuery) -> None:
        doc_id=int(cb.data.rsplit(":",1)[1]); row=db.fetch_one("SELECT d.*,o.name org FROM portal_documents d JOIN organizations o ON o.id=d.organization_id WHERE d.id=?",(doc_id,))
        if not row or not _can_read_portal(cb.from_user.id,row["org"]): return await cb.answer("Немає доступу.",show_alert=True)
        if roles_for(cb.from_user.id,row["org"]) & {"admin","fop"} or (row["published"] and ((row["for_user"] and "user" in roles_for(cb.from_user.id,row["org"])) or (row["for_leadership"] and "leadership" in roles_for(cb.from_user.id,row["org"])) or "accountant" in roles_for(cb.from_user.id,row["org"]))):
            await bot.send_document(cb.message.chat.id,row["file_id"],caption=row["name"])
        else: return await cb.answer("Документ не опублікований для вашої ролі.",show_alert=True)
        await cb.answer()

    @dp.message(F.text == "📇 Контакти")
    async def contacts_menu(m: types.Message) -> None:
        roles=_user_roles(m.from_user.id); required={"admin","fop"} if roles & {"admin","fop"} else {"user","leadership","accountant"}
        await m.answer("Оберіть організацію:",reply_markup=_org_keyboard(m.from_user.id,"contactmenu",required))

    @dp.callback_query(F.data.startswith("p:contactmenu:"))
    async def contacts_by_org(cb: types.CallbackQuery) -> None:
        org_id=int(cb.data.rsplit(":",1)[1]); org=_org(org_id)
        if not org or not _can_read_portal(cb.from_user.id,org["name"]): return await cb.answer("Немає доступу.",show_alert=True)
        rows=db.fetch_all("SELECT id,name,phone,description FROM contacts WHERE organization_id=? AND active=1 ORDER BY name",(org_id,))
        for r in rows: await cb.message.answer(f"{r['name']} · {r['phone']}\n{r['description']}")
        if roles_for(cb.from_user.id,org["name"]) & {"admin","fop"}:
            kb=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="➕ Додати контакт",callback_data=f"p:addcontact:{org_id}")]])
            await cb.message.answer("Керування контактами:",reply_markup=kb)
        elif not rows: await cb.message.answer("Контактів поки немає.")
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:addcontact:"))
    async def contact_add_start(cb: types.CallbackQuery,state: FSMContext) -> None:
        org_id=int(cb.data.rsplit(":",1)[1]); org=_org(org_id)
        if not org or not _is_manager(cb.from_user.id,org["name"]): return await cb.answer("Немає доступу.",show_alert=True)
        await state.update_data(org_id=org_id); await state.set_state(PortalForm.contact_value)
        await cb.message.answer("Введіть: ПІБ | телефон | опис контакту"); await cb.answer()

    @dp.message(PortalForm.contact_value)
    async def contact_add_save(m: types.Message,state: FSMContext) -> None:
        data=await state.get_data(); org_id=int(data.get("org_id") or 0); org=_org(org_id)
        if not org or not _is_manager(m.from_user.id,org["name"]): await state.clear(); return await m.answer("Немає доступу.")
        try: name,phone,desc=[part.strip() for part in (m.text or "").split("|",2)]
        except ValueError: return await m.answer("Формат: ПІБ | телефон | опис контакту")
        if not name or not phone: return await m.answer("Вкажіть ПІБ і телефон.")
        contact_id=db.execute("INSERT INTO contacts(organization_id,name,phone,description) VALUES(?,?,?,?)",(org_id,name,phone,desc))
        _audit(m.from_user.id,"add_contact","contact",contact_id,org["name"])
        await state.clear(); await m.answer("Контакт додано.")

    @dp.message(Command("contact"))
    async def add_contact(m: types.Message) -> None:
        parts=(m.text or "").split(maxsplit=1)
        if len(parts)<2 or "|" not in parts[1]: return await m.answer("Формат: /contact ID_ОРГАНІЗАЦІЇ | ПІБ | телефон | опис")
        try: org_raw, name, phone, desc=[p.strip() for p in parts[1].split("|",3)]; org=_org(int(org_raw))
        except Exception: return await m.answer("Перевірте формат команди.")
        if not org or not _is_manager(m.from_user.id,org["name"]): return await m.answer("Немає доступу до цієї організації.")
        contact_id=db.execute("INSERT INTO contacts(organization_id,name,phone,description) VALUES(?,?,?,?)",(int(org_raw),name,phone,desc))
        _audit(m.from_user.id,"add_contact","contact",contact_id,org["name"]); await m.answer("Контакт додано.")

    @dp.callback_query(F.data.startswith("p:sa:"))
    async def superadmin_router(cb: types.CallbackQuery,state: FSMContext) -> None:
        if not is_superadmin(cb.from_user.id): return await cb.answer("Немає доступу.",show_alert=True)
        action=cb.data.split(":",2)[2]
        if action=="branches":
            branches=db.fetch_all("SELECT * FROM branches ORDER BY id")
            rows=[[InlineKeyboardButton(text=f"{b['name']} · {b['status']} · {b['payment_status']}",callback_data=f"p:sabranch:{b['id']}")] for b in branches]
            rows.append([InlineKeyboardButton(text="➕ Додати гілку",callback_data="p:sacreatebranch")])
            await cb.message.answer("Гілки:",reply_markup=InlineKeyboardMarkup(inline_keyboard=rows))
        elif action=="orgs":
            orgs=db.fetch_all("SELECT o.*,b.name branch FROM organizations o JOIN branches b ON b.id=o.branch_id ORDER BY o.name")
            rows=[[InlineKeyboardButton(text=f"{o['name']} · {o['kind']} · {o['branch']}",callback_data=f"p:saorg:{o['id']}")] for o in orgs]
            rows.append([InlineKeyboardButton(text="➕ Додати організацію",callback_data="p:sacreateorg")])
            await cb.message.answer("ОСББ, кооперативи та ФОП:",reply_markup=InlineKeyboardMarkup(inline_keyboard=rows))
        elif action=="users":
            await cb.message.answer("Керування ролями. Формат:\n/role ID | Назва організації | admin/accountant/viewer/leadership/fop | keywords (лише для viewer)\nВідкликати: /roleoff ID | Назва організації | роль")
        elif action=="audit":
            rows=db.fetch_all("SELECT user_id,action,entity_type,entity_id,osbb,details,created_at FROM audit_log ORDER BY id DESC LIMIT 30")
            text="🧾 Останні записи аудиту:\n"+("\n".join(f"{r['created_at']} · {r['user_id']} · {r['action']} · {r['osbb'] or r['entity_type']} · {r['details'] or ''}" for r in rows) if rows else "Журнал порожній.")
            await cb.message.answer(text)
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:sabranch:"))
    async def branch_settings(cb: types.CallbackQuery) -> None:
        if not is_superadmin(cb.from_user.id): return await cb.answer("Немає доступу.",show_alert=True)
        branch_id=int(cb.data.rsplit(":",1)[1]); b=db.fetch_one("SELECT * FROM branches WHERE id=?",(branch_id,))
        if not b: return await cb.answer("Гілку не знайдено.",show_alert=True)
        kb=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text=f"Оплата: {b['payment_status']} · змінити",callback_data=f"p:sapayment:{branch_id}")],
            [InlineKeyboardButton(text=f"Стан: {b['status']} · змінити",callback_data=f"p:sastatus:{branch_id}")],
            [InlineKeyboardButton(text="⬅️ До гілок",callback_data="p:sa:branches")],
        ])
        await cb.message.answer(f"{b['name']} · власник: {b['owner_type']} {b['owner_user_id'] or ''} · оплачено до {b['paid_until'] or 'без дати'}",reply_markup=kb); await cb.answer()

    @dp.callback_query(F.data.startswith("p:sapayment:"))
    async def branch_payment_menu(cb: types.CallbackQuery) -> None:
        if not is_superadmin(cb.from_user.id): return await cb.answer("Немає доступу.",show_alert=True)
        branch_id=int(cb.data.rsplit(":",1)[1]); statuses=["paid","unpaid","trial","blocked"]
        kb=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text=s,callback_data=f"p:setpayment:{branch_id}:{s}")] for s in statuses]+[[InlineKeyboardButton(text="📅 Встановити дату оплати / trial",callback_data=f"p:paiduntil:{branch_id}")]])
        await cb.message.answer("Виберіть статус оплати:",reply_markup=kb); await cb.answer()

    @dp.callback_query(F.data.startswith("p:setpayment:"))
    async def branch_payment_set(cb: types.CallbackQuery) -> None:
        if not is_superadmin(cb.from_user.id): return await cb.answer("Немає доступу.",show_alert=True)
        _,_,branch_raw,status=cb.data.split(":",3)
        if status not in {"paid","unpaid","trial","blocked"}: return await cb.answer("Статус невірний.",show_alert=True)
        bid=int(branch_raw); db.execute("UPDATE branches SET payment_status=? WHERE id=?",(status,bid)); _audit(cb.from_user.id,"set_branch_payment","branch",bid,details=status)
        await cb.answer("Статус оплати збережено.",show_alert=True)

    @dp.callback_query(F.data.startswith("p:paiduntil:"))
    async def branch_paid_until_start(cb: types.CallbackQuery,state: FSMContext) -> None:
        if not is_superadmin(cb.from_user.id): return await cb.answer("Немає доступу.",show_alert=True)
        branch_id=int(cb.data.rsplit(":",1)[1]); await state.update_data(branch_id=branch_id); await state.set_state(PortalForm.payment_change)
        await cb.message.answer("Введіть дату YYYY-MM-DD до якої діє оплата/trial або «-» без дати."); await cb.answer()

    @dp.message(PortalForm.payment_change)
    async def branch_paid_until_save(m: types.Message,state: FSMContext) -> None:
        if not is_superadmin(m.from_user.id): await state.clear(); return await m.answer("Немає доступу.")
        data=await state.get_data(); raw=(m.text or "").strip()
        if raw!="-":
            try: raw=date.fromisoformat(raw).isoformat()
            except ValueError: return await m.answer("Введіть YYYY-MM-DD або «-».")
        bid=int(data.get("branch_id") or 0); db.execute("UPDATE branches SET paid_until=? WHERE id=?",(None if raw=="-" else raw,bid))
        _audit(m.from_user.id,"set_paid_until","branch",bid,details=raw)
        await state.clear(); await m.answer("Дату оплати / trial збережено.")

    @dp.callback_query(F.data.startswith("p:sastatus:"))
    async def branch_status_toggle(cb: types.CallbackQuery) -> None:
        if not is_superadmin(cb.from_user.id): return await cb.answer("Немає доступу.",show_alert=True)
        bid=int(cb.data.rsplit(":",1)[1]); row=db.fetch_one("SELECT status FROM branches WHERE id=?",(bid,))
        if not row: return await cb.answer("Гілку не знайдено.",show_alert=True)
        new="paused" if row["status"]=="active" else "active"; db.execute("UPDATE branches SET status=? WHERE id=?",(new,bid)); _audit(cb.from_user.id,"set_branch_status","branch",bid,details=new)
        await cb.answer(f"Стан гілки: {new}.",show_alert=True)

    @dp.callback_query(F.data == "p:sacreatebranch")
    async def branch_create_start(cb: types.CallbackQuery,state: FSMContext) -> None:
        if not is_superadmin(cb.from_user.id): return await cb.answer("Немає доступу.",show_alert=True)
        await state.set_state(PortalForm.branch_create); await cb.message.answer("Введіть: назва гілки | admin або fop | Telegram ID власника (ID можна лишити порожнім).")
        await cb.answer()

    @dp.message(PortalForm.branch_create)
    async def branch_create_save(m: types.Message,state: FSMContext) -> None:
        if not is_superadmin(m.from_user.id): await state.clear(); return await m.answer("Немає доступу.")
        try:
            name,owner_type,owner_raw=[p.strip() for p in (m.text or "").split("|",2)]
            if owner_type not in {"admin","fop"}: raise ValueError
            owner_id=int(owner_raw) if owner_raw else None
        except Exception: return await m.answer("Формат: назва | admin або fop | Telegram ID власника")
        try: branch_id=db.execute("INSERT INTO branches(name,owner_type,owner_user_id) VALUES(?,?,?)",(name,owner_type,owner_id))
        except Exception: return await m.answer("Не вдалося додати гілку: можливо, назва вже використовується.")
        if owner_id:
            db.execute("INSERT OR IGNORE INTO users(telegram_id,full_name) VALUES(?,?)",(owner_id,owner_type.upper()))
        _audit(m.from_user.id,"create_branch","branch",branch_id,details=f"owner={owner_type}:{owner_id}")
        await state.clear(); await m.answer(f"Гілку «{name}» створено. Тепер додайте організації у меню /admin → ОСББ та кооперативи.")

    @dp.callback_query(F.data == "p:sacreateorg")
    async def org_create_start(cb: types.CallbackQuery,state: FSMContext) -> None:
        if not is_superadmin(cb.from_user.id): return await cb.answer("Немає доступу.",show_alert=True)
        branches=db.fetch_all("SELECT id,name FROM branches WHERE status='active' ORDER BY id")
        await state.set_state(PortalForm.org_create)
        await cb.message.answer("Введіть: ID гілки | назва | osbb або cooperative або fop\nГілки: "+", ".join(f"{b['id']}={b['name']}" for b in branches)); await cb.answer()

    @dp.message(PortalForm.org_create)
    async def org_create_save(m: types.Message,state: FSMContext) -> None:
        if not is_superadmin(m.from_user.id): await state.clear(); return await m.answer("Немає доступу.")
        try:
            branch_raw,name,kind=[p.strip() for p in (m.text or "").split("|",2)]; branch_id=int(branch_raw)
            if kind not in {"osbb","cooperative","fop"} or not db.fetch_one("SELECT 1 FROM branches WHERE id=?",(branch_id,)) or db.fetch_one("SELECT 1 FROM organizations WHERE name=?",(name,)): raise ValueError
        except Exception: return await m.answer("Формат: ID гілки | назва | osbb/cooperative/fop")
        org_id=db.execute("INSERT INTO organizations(branch_id,name,kind,resident_code,leadership_code) VALUES(?,?,?,?,?)",(branch_id,name,kind,db.create_join_code(),db.create_join_code()))
        branch=db.fetch_one("SELECT owner_user_id,owner_type FROM branches WHERE id=?",(branch_id,))
        if branch and branch["owner_user_id"]:
            role="fop" if branch["owner_type"]=="fop" else "admin"
            db.execute("INSERT OR IGNORE INTO memberships(telegram_id,organization_id,role) VALUES(?,?,?)",(branch["owner_user_id"],org_id,role))
        _audit(m.from_user.id,"create_organization","organization",org_id,name,kind)
        await state.clear(); await m.answer(f"Створено {name} ({kind}). Коди для мешканців і Leadership: /admin → організація → коди.")

    @dp.callback_query(F.data.startswith("p:saorg:"))
    async def org_admin_view(cb: types.CallbackQuery) -> None:
        if not is_superadmin(cb.from_user.id): return await cb.answer("Немає доступу.",show_alert=True)
        org_id=int(cb.data.rsplit(":",1)[1]); org=_org(org_id)
        if not org: return await cb.answer("Організацію не знайдено.",show_alert=True)
        kb=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="Коди доступу",callback_data=f"p:codes:{org_id}")],
            [InlineKeyboardButton(text="Відкрити налаштування будинку",callback_data=f"p:building:{org_id}")],
            [InlineKeyboardButton(text="⬅️ До організацій",callback_data="p:sa:orgs")],
        ])
        await cb.message.answer(f"{org['name']} · {org['kind']} · гілка {org['branch_id']}",reply_markup=kb); await cb.answer()

    @dp.message(Command("role"))
    async def grant_role(m: types.Message) -> None:
        if not is_superadmin(m.from_user.id): return await m.answer("Команда доступна лише Superadmin.")
        parts=(m.text or "").split(maxsplit=1)
        if len(parts)<2: return await m.answer("Формат: /role ID | Назва організації | роль | keywords")
        try:
            uid_raw,name,role,*rest=[p.strip() for p in parts[1].split("|")]; uid=int(uid_raw); org=_org_by_name(name)
            if not org or role not in {"admin","accountant","viewer","leadership","fop"}: raise ValueError
        except Exception: return await m.answer("Перевірте ID, назву організації та роль.")
        keywords=rest[0] if rest else ""
        if role=="viewer" and not keywords: return await m.answer("Для Viewer вкажіть ключові слова через кому.")
        _upsert_user(types.User(id=uid,is_bot=False,first_name="Користувач"))
        db.execute("INSERT INTO memberships(telegram_id,organization_id,role,keywords) VALUES(?,?,?,?) ON CONFLICT(telegram_id,organization_id,role) DO UPDATE SET keywords=excluded.keywords,active=1",(uid,org["id"],role,keywords))
        _audit(m.from_user.id,"grant_role","membership",None,name,f"{uid} {role}")
        await m.answer(f"Роль {role} призначено користувачу {uid} в {name}.")

    @dp.message(Command("roleoff"))
    async def revoke_role(m: types.Message) -> None:
        if not is_superadmin(m.from_user.id): return await m.answer("Команда доступна лише Superadmin.")
        try:
            parts=(m.text or "").split(maxsplit=1); uid_raw,name,role=[p.strip() for p in parts[1].split("|")]; org=_org_by_name(name); uid=int(uid_raw)
            if not org: raise ValueError
        except Exception: return await m.answer("Формат: /roleoff ID | Назва організації | роль")
        db.execute("UPDATE memberships SET active=0 WHERE telegram_id=? AND organization_id=? AND role=?",(uid,org["id"],role))
        _audit(m.from_user.id,"revoke_role","membership",None,name,f"{uid} {role}")
        await m.answer(f"Доступ відкликано: {uid} · {role} · {name}.")

    @dp.callback_query(F.data.startswith("p:export:"))
    async def export_readings_prompt(cb: types.CallbackQuery, state: FSMContext) -> None:
        org_id=int(cb.data.rsplit(":",1)[1]); org=_org(org_id)
        if not org or not _can_export_keys(cb.from_user.id,org["name"]): return await cb.answer("Немає доступу.",show_alert=True)
        await state.update_data(org_id=org_id)
        await state.set_state(PortalForm.export_year)
        await cb.message.answer(f"Введіть рік для {org['name']} (наприклад, 2026):")
        await cb.answer()

    @dp.message(PortalForm.export_year)
    async def export_readings_year(m: types.Message, state: FSMContext) -> None:
        data=await state.get_data(); org_id=int(data.get("org_id") or 0); org=_org(org_id)
        try:
            year=int((m.text or "").strip())
            if year < 2000 or year > 2100: raise ValueError
        except ValueError: return await m.answer("Введіть чотиризначний рік від 2000 до 2100.")
        if not org or not _can_export_indicators(m.from_user.id,org["name"]):
            await state.clear(); return await m.answer("Немає доступу.")
        payload=_build_readings_xlsx(org_id,year)
        await state.clear()
        await m.answer_document(BufferedInputFile(payload,filename=f"Показники_{org['name']}_{year}.xlsx"))

    @dp.callback_query(F.data.startswith("p:pending:"))
    async def pending_readings(cb: types.CallbackQuery) -> None:
        org_id=int(cb.data.rsplit(":",1)[1]); org=_org(org_id)
        if not org or not _is_admin(cb.from_user.id,org["name"]): return await cb.answer("Немає доступу.",show_alert=True)
        rows=db.fetch_all("SELECT r.id,r.metric,r.period,r.value,r.photo_file_id,a.number FROM readings r JOIN apartments a ON a.id=r.apartment_id WHERE a.organization_id=? AND r.status='pending' ORDER BY r.created_at",(org_id,))
        if not rows: await cb.message.answer("Показників на перевірку немає.")
        for r in rows[:50]:
            text=f"Квартира {r['number']} · {r['period']}\n{METRICS.get(r['metric'],r['metric'])}: {r['value']}"
            kb=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="✅ Прийняти",callback_data=f"p:readingaccept:{r['id']}")],[InlineKeyboardButton(text="❌ Відхилити",callback_data=f"p:readingreject:{r['id']}")]])
            if r["photo_file_id"]: await bot.send_photo(cb.message.chat.id,r["photo_file_id"],caption=text,reply_markup=kb)
            else: await cb.message.answer(text,reply_markup=kb)
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:readingaccept:"))
    async def reading_accept(cb: types.CallbackQuery) -> None:
        reading_id=int(cb.data.rsplit(":",1)[1]); row=db.fetch_one("SELECT r.*,o.name org FROM readings r JOIN apartments a ON a.id=r.apartment_id JOIN organizations o ON o.id=a.organization_id WHERE r.id=?",(reading_id,))
        if not row or not _is_admin(cb.from_user.id,row["org"]): return await cb.answer("Немає доступу.",show_alert=True)
        db.execute("UPDATE readings SET status='accepted',photo_file_id=NULL,decided_by=?,decided_at=?,decision_note='' WHERE id=?",(cb.from_user.id,db.utc_now(),reading_id))
        _audit(cb.from_user.id,"accept_reading","reading",reading_id,row["org"])
        await cb.message.edit_caption(caption=(cb.message.caption or "")+"\n✅ Прийнято; фото прибрано з обліку.",reply_markup=None) if cb.message.photo else cb.message.edit_text((cb.message.text or "")+"\n✅ Прийнято; фото прибрано з обліку.",reply_markup=None)
        try: await bot.send_message(row["submitted_by"],f"Показник {METRICS.get(row['metric'],row['metric'])} за {row['period']} прийнято.")
        except Exception: logger.info("Could not notify resident about accepted reading id=%s",reading_id)
        await cb.answer("Показник прийнято.")

    @dp.callback_query(F.data.startswith("p:readingreject:"))
    async def reading_reject_start(cb: types.CallbackQuery, state: FSMContext) -> None:
        reading_id=int(cb.data.rsplit(":",1)[1]); row=db.fetch_one("SELECT r.id,o.name org FROM readings r JOIN apartments a ON a.id=r.apartment_id JOIN organizations o ON o.id=a.organization_id WHERE r.id=?",(reading_id,))
        if not row or not _is_admin(cb.from_user.id,row["org"]): return await cb.answer("Немає доступу.",show_alert=True)
        await state.update_data(reading_id=reading_id); await state.set_state(PortalForm.reading_review_note)
        await cb.message.answer("Напишіть мешканцю причину відхилення показника:"); await cb.answer()

    @dp.message(PortalForm.reading_review_note)
    async def reading_reject_save(m: types.Message,state: FSMContext) -> None:
        data=await state.get_data(); rid=int(data.get("reading_id") or 0); note=(m.text or "").strip()
        row=db.fetch_one("SELECT r.submitted_by,r.metric,r.period,o.name org FROM readings r JOIN apartments a ON a.id=r.apartment_id JOIN organizations o ON o.id=a.organization_id WHERE r.id=?",(rid,))
        if not note or not row or not _is_admin(m.from_user.id,row["org"]): return await m.answer("Вкажіть причину відхилення.")
        db.execute("UPDATE readings SET status='rejected',decision_note=?,photo_file_id=NULL,decided_by=?,decided_at=? WHERE id=?",(note,m.from_user.id,db.utc_now(),rid))
        _audit(m.from_user.id,"reject_reading","reading",rid,row["org"],note)
        try: await bot.send_message(row["submitted_by"],f"Показник {METRICS.get(row['metric'],row['metric'])} за {row['period']} відхилено. Причина: {note}\nМожна надіслати його повторно.")
        except Exception: logger.info("Could not notify resident about rejected reading id=%s",rid)
        await state.clear(); await m.answer("Показник відхилено, пояснення надіслано мешканцю; фото прибрано з обліку.")

    @dp.callback_query(F.data.startswith("p:contacts:"))
    async def contacts_view(cb: types.CallbackQuery) -> None:
        org_id = int(cb.data.rsplit(":", 1)[1])
        org = _org(org_id)
        if not org or not _can_read_portal(cb.from_user.id, org["name"]):
            return await cb.answer("Немає доступу.", show_alert=True)
        rows = db.fetch_all("SELECT name,phone,description FROM contacts WHERE organization_id=? AND active=1 ORDER BY name", (org_id,))
        text = "📇 Контакти\n" + ("\n\n".join(f"{r['name']} · {r['phone']}\n{r['description']}" for r in rows) if rows else "Контактів поки немає.")
        await cb.message.answer(text)
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:jobs:"))
    async def resident_jobs(cb: types.CallbackQuery) -> None:
        org_id = int(cb.data.rsplit(":", 1)[1])
        org = _org(org_id)
        if not org or not _can_read_portal(cb.from_user.id, org["name"]):
            return await cb.answer("Немає доступу.", show_alert=True)
        rows = db.fetch_all("SELECT task_text,status,priority,deadline,stages,comments FROM jobs WHERE osbb=? ORDER BY id DESC LIMIT 100", (org["name"],))
        if not rows:
            await cb.message.answer("Робіт поки немає.")
        else:
            for r in rows:
                await cb.message.answer(f"{r['status']} · {r['priority']}\n{r['task_text']}\nДедлайн: {r['deadline'] or '—'}\nЕтапи: {r['stages'] or '—'}\nКоментарі: {r['comments'] or '—'}")
        await cb.answer()

    @dp.message(F.text == "🏘️ Мешканці та будинок")
    async def admin_building(m: types.Message) -> None:
        await m.answer("Оберіть організацію для керування будинком:", reply_markup=_org_keyboard(m.from_user.id, "building", {"admin"}))

    @dp.callback_query(F.data.startswith("p:building:"))
    async def admin_building_org(cb: types.CallbackQuery) -> None:
        org_id = int(cb.data.rsplit(":", 1)[1])
        org = _org(org_id)
        if not org or not _is_admin(cb.from_user.id, org["name"]):
            return await cb.answer("Немає права керувати будинком.", show_alert=True)
        kb = InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🏠 Кількість квартир і нумерація", callback_data=f"p:aptcount:{org_id}")],
            [InlineKeyboardButton(text="⚙️ Показники по квартирах", callback_data=f"p:meterlist:{org_id}")],
            [InlineKeyboardButton(text=f"⏸ Подача всіх показників: {'увімкнена' if org['meters_enabled'] else 'зупинена'}", callback_data=f"p:orgmeters:{org_id}")],
            [InlineKeyboardButton(text="🔑 Коди мешканців і Leadership", callback_data=f"p:codes:{org_id}")],
            [InlineKeyboardButton(text="📥 Вивантажити ключі квартир", callback_data=f"p:pinexport:{org_id}")],
            [InlineKeyboardButton(text="⬅️ Назад", callback_data="p:home")],
        ])
        await cb.message.answer(f"Керування {org['name']} · квартир зараз: {org['apartment_count']}", reply_markup=kb)
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:orgmeters:"))
    async def toggle_all_meters(cb: types.CallbackQuery) -> None:
        org_id=int(cb.data.rsplit(":",1)[1]); org=_org(org_id)
        if not org or not _is_admin(cb.from_user.id,org["name"]): return await cb.answer("Немає доступу.",show_alert=True)
        new=0 if org["meters_enabled"] else 1
        db.execute("UPDATE organizations SET meters_enabled=? WHERE id=?",(new,org_id))
        _audit(cb.from_user.id,"toggle_all_meters","organization",org_id,org["name"],f"enabled={new}")
        await cb.answer("Подачу показників зупинено." if not new else "Подачу показників відновлено.",show_alert=True)

    @dp.callback_query(F.data.startswith("p:aptcount:"))
    async def apartment_count_start(cb: types.CallbackQuery, state: FSMContext) -> None:
        org_id = int(cb.data.rsplit(":", 1)[1])
        org = _org(org_id)
        if not org or not _is_admin(cb.from_user.id, org["name"]):
            return await cb.answer("Немає доступу.", show_alert=True)
        await state.update_data(org_id=org_id)
        await state.set_state(PortalForm.apartment_count)
        await cb.message.answer("Введіть загальну кількість квартир. Наявні квартири й історія показників не видаляються.")
        await cb.answer()

    @dp.message(PortalForm.apartment_count)
    async def apartment_count_save(m: types.Message, state: FSMContext) -> None:
        data = await state.get_data()
        org_id = int(data.get("org_id") or 0)
        org = _org(org_id)
        if not org or not _is_admin(m.from_user.id, org["name"]):
            await state.clear()
            return await m.answer("Немає доступу до цього ОСББ.")
        try:
            count = int((m.text or "").strip())
            if count < 1 or count > 5000:
                raise ValueError
        except ValueError:
            return await m.answer("Вкажіть ціле число від 1 до 5000.")
        present = {r["number"] for r in db.fetch_all("SELECT number FROM apartments WHERE organization_id=?", (org_id,))}
        previous=int(org["apartment_count"] or 0)
        with db.transaction() as conn:
            conn.execute("UPDATE organizations SET apartment_count=? WHERE id=?", (count, org_id))
            for i in range(previous + 1, count + 1):
                number = str(i)
                if number not in present:
                    conn.execute("INSERT INTO apartments(organization_id,number) VALUES(?,?)", (org_id, number))
        _audit(m.from_user.id, "set_apartment_count", "organization", org_id, org["name"], f"count={count}")
        await state.clear()
        note=" Наявні квартири не деактивовано; за потреби заархівуйте зайві у картці квартири." if count<previous else ""
        await m.answer(f"Збережено кількість квартир: {count}. Нові квартири додано з наступних вільних номерів. Нумерацію можна змінити в картці квартири.{note}")

    @dp.callback_query(F.data.startswith("p:meterlist:"))
    async def apartment_list(cb: types.CallbackQuery) -> None:
        org_id = int(cb.data.rsplit(":", 1)[1])
        org = _org(org_id)
        if not org or not _is_admin(cb.from_user.id, org["name"]):
            return await cb.answer("Немає доступу.", show_alert=True)
        rows = db.fetch_all("SELECT id,number FROM apartments WHERE organization_id=? AND active=1 ORDER BY number", (org_id,))
        buttons = [[InlineKeyboardButton(text=f"Квартира {r['number']}", callback_data=f"p:meterapt:{r['id']}")] for r in rows[:90]]
        buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data=f"p:building:{org_id}")])
        await cb.message.answer("Оберіть квартиру для налаштування:", reply_markup=InlineKeyboardMarkup(inline_keyboard=buttons))
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:meterapt:"))
    async def meter_apartment(cb: types.CallbackQuery) -> None:
        apt_id = int(cb.data.rsplit(":", 1)[1])
        apt = db.fetch_one("SELECT a.*,o.name org,o.id org_id FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?", (apt_id,))
        if not apt or not _is_admin(cb.from_user.id, apt["org"]):
            return await cb.answer("Немає доступу.", show_alert=True)
        settings = {r["metric"]: r for r in db.fetch_all("SELECT * FROM meter_settings WHERE apartment_id=?", (apt_id,))}
        rows = []
        for metric, title in METRICS.items():
            s = settings.get(metric, {"enabled": 0, "mode": "meter"})
            if metric == "people":
                continue
            mark = "✅" if s["enabled"] else "⛔"
            rows.append([InlineKeyboardButton(text=f"{mark} {title} · {'особи' if s['mode']=='people' else 'лічильник'}", callback_data=f"p:togglemetric:{apt_id}:{metric}")])
        rows += [
            [InlineKeyboardButton(text="👥 Задати кількість осіб / період", callback_data=f"p:peopleperiod:{apt_id}")],
            [InlineKeyboardButton(text="⚡ Лічильник / особи для електроенергії", callback_data=f"p:electricitymode:{apt_id}")],
            [InlineKeyboardButton(text="📄 Документ у кабінет квартири", callback_data=f"p:aptdoc:{apt_id}")],
            [InlineKeyboardButton(text="✏️ Змінити номер квартири", callback_data=f"p:rename:{apt_id}")],
            [InlineKeyboardButton(text="🗄 Архівувати квартиру", callback_data=f"p:archiveapt:{apt_id}")],
            [InlineKeyboardButton(text="📝 Інформація для мешканця", callback_data=f"p:aptinfo:{apt_id}")],
            [InlineKeyboardButton(text="⬅️ До квартир", callback_data=f"p:meterlist:{apt['org_id']}")],
        ]
        await cb.message.answer(f"Квартира {apt['number']} · {apt['org']}", reply_markup=InlineKeyboardMarkup(inline_keyboard=rows))
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:archiveapt:"))
    async def archive_apartment(cb: types.CallbackQuery) -> None:
        apt_id=int(cb.data.rsplit(":",1)[1]); apt=db.fetch_one("SELECT a.organization_id,a.number,o.name org FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?",(apt_id,))
        if not apt or not _is_admin(cb.from_user.id,apt["org"]): return await cb.answer("Немає доступу.",show_alert=True)
        db.execute("UPDATE apartments SET active=0 WHERE id=?",(apt_id,))
        active=db.fetch_one("SELECT count(*) n FROM apartments WHERE organization_id=? AND active=1",(apt["organization_id"],))["n"]
        db.execute("UPDATE organizations SET apartment_count=? WHERE id=?",(active,apt["organization_id"]))
        _audit(cb.from_user.id,"archive_apartment","apartment",apt_id,apt["org"],apt["number"])
        await cb.answer(f"Квартиру {apt['number']} заархівовано; показники та документи збережено.",show_alert=True)

    @dp.callback_query(F.data.startswith("p:togglemetric:"))
    async def toggle_metric(cb: types.CallbackQuery) -> None:
        _, _, apt_raw, metric = cb.data.split(":", 3)
        apt_id = int(apt_raw)
        apt = db.fetch_one("SELECT a.number,o.name org FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?", (apt_id,))
        if not apt or not _is_admin(cb.from_user.id, apt["org"]):
            return await cb.answer("Немає доступу.", show_alert=True)
        row = db.fetch_one("SELECT enabled,mode FROM meter_settings WHERE apartment_id=? AND metric=?", (apt_id, metric))
        new_enabled = 0 if row and row["enabled"] else 1
        mode = row["mode"] if row else "meter"
        db.execute("INSERT INTO meter_settings(apartment_id,metric,mode,enabled) VALUES(?,?,?,?) ON CONFLICT(apartment_id,metric) DO UPDATE SET enabled=excluded.enabled", (apt_id, metric, mode, new_enabled))
        _audit(cb.from_user.id, "toggle_meter", "apartment", apt_id, apt["org"], f"metric={metric}; enabled={new_enabled}")
        await cb.answer("Налаштування збережено.")

    @dp.callback_query(F.data.startswith("p:electricitymode:"))
    async def electricity_mode(cb: types.CallbackQuery) -> None:
        apt_id=int(cb.data.rsplit(":",1)[1]); apt=db.fetch_one("SELECT a.number,o.name org FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?",(apt_id,))
        if not apt or not _is_admin(cb.from_user.id,apt["org"]): return await cb.answer("Немає доступу.",show_alert=True)
        row=db.fetch_one("SELECT mode,enabled FROM meter_settings WHERE apartment_id=? AND metric='electricity'",(apt_id,))
        mode="people" if row and row["mode"]=="meter" else "meter"
        enabled=int(row["enabled"]) if row else 0
        db.execute("INSERT INTO meter_settings(apartment_id,metric,mode,enabled) VALUES(?,'electricity',?,?) ON CONFLICT(apartment_id,metric) DO UPDATE SET mode=excluded.mode",(apt_id,mode,enabled))
        _audit(cb.from_user.id,"change_electricity_mode","apartment",apt_id,apt["org"],mode)
        await cb.answer("Режим змінено: "+("кількість осіб" if mode=="people" else "лічильник"),show_alert=True)

    @dp.callback_query(F.data.startswith("p:peopleperiod:"))
    async def people_period_start(cb: types.CallbackQuery,state: FSMContext) -> None:
        apt_id=int(cb.data.rsplit(":",1)[1]); apt=db.fetch_one("SELECT a.number,o.name org FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?",(apt_id,))
        if not apt or not _is_admin(cb.from_user.id,apt["org"]): return await cb.answer("Немає доступу.",show_alert=True)
        await state.update_data(apartment_id=apt_id); await state.set_state(PortalForm.resident_period)
        await cb.message.answer(f"Для квартири {apt['number']} введіть кількість і період за актом фактичного проживання у форматі:\n4, 2026-07-01, 2026-09-30")
        await cb.answer()

    @dp.message(PortalForm.resident_period)
    async def people_period_save(m: types.Message,state: FSMContext) -> None:
        data=await state.get_data(); apt_id=int(data.get("apartment_id") or 0)
        apt=db.fetch_one("SELECT a.number,a.organization_id,o.name org FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?",(apt_id,))
        if not apt or not _is_admin(m.from_user.id,apt["org"]): await state.clear(); return await m.answer("Немає доступу.")
        try:
            count_s,start_s,end_s=[part.strip() for part in (m.text or "").split(",")]
            count=int(count_s); start=date.fromisoformat(start_s); end=date.fromisoformat(end_s)
            if count<0 or count>50 or end<start: raise ValueError
        except Exception: return await m.answer("Перевірте формат, кількість (0–50) та дати: число, YYYY-MM-DD, YYYY-MM-DD.")
        await state.update_data(resident_count=count,period_start=start.isoformat(),period_end=end.isoformat())
        await state.set_state(PortalForm.resident_period_act)
        await m.answer("Тепер завантажте підписаний Акт фактичного проживання (фото або PDF). Він з’явиться в кабінеті цієї квартири.")

    @dp.message(PortalForm.resident_period_act)
    async def people_period_act_save(m: types.Message,state: FSMContext) -> None:
        data=await state.get_data(); apt_id=int(data.get("apartment_id") or 0)
        apt=db.fetch_one("SELECT a.number,a.organization_id,o.name org FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?",(apt_id,))
        if not apt or not _is_admin(m.from_user.id,apt["org"]): await state.clear(); return await m.answer("Немає доступу.")
        if m.photo:
            file_id=m.photo[-1].file_id; title=f"Акт фактичного проживання · {data['period_start']} — {data['period_end']}"
        elif m.document and (m.document.mime_type=="application/pdf" or (m.document.file_name or "").lower().endswith(".pdf")):
            file_id=m.document.file_id; title=(m.document.file_name or "Акт фактичного проживання")[:150]
        else:
            return await m.answer("Надішліть фото акту або PDF-файл.")
        with db.transaction() as conn:
            conn.execute("INSERT INTO resident_periods(apartment_id,person_count,starts_on,ends_on,act_file_id,created_by) VALUES(?,?,?,?,?,?)",(apt_id,data["resident_count"],data["period_start"],data["period_end"],file_id,m.from_user.id))
            today=datetime.now(LOCAL_TZ).date().isoformat()
            if data["period_start"]<=today<=data["period_end"]:
                conn.execute("UPDATE apartments SET resident_count=? WHERE id=?",(data["resident_count"],apt_id))
            conn.execute("INSERT INTO portal_documents(organization_id,apartment_id,name,file_id,published,for_user,created_by) VALUES(?,?,?,?,1,1,?)",(apt["organization_id"],apt_id,title,file_id,m.from_user.id))
        _audit(m.from_user.id,"set_resident_period","apartment",apt_id,apt["org"],f"{data['resident_count']}; {data['period_start']}; {data['period_end']}; act attached")
        await state.clear(); await m.answer(f"Збережено: {data['resident_count']} осіб з {data['period_start']} до {data['period_end']}. Акт додано в кабінет квартири та річний Excel.")

    @dp.callback_query(F.data.startswith("p:rename:"))
    async def rename_apartment_start(cb: types.CallbackQuery, state: FSMContext) -> None:
        apt_id = int(cb.data.rsplit(":", 1)[1])
        apt = db.fetch_one("SELECT a.number,o.name org FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?", (apt_id,))
        if not apt or not _is_admin(cb.from_user.id, apt["org"]):
            return await cb.answer("Немає доступу.", show_alert=True)
        await state.update_data(apartment_id=apt_id)
        await state.set_state(PortalForm.apartment_rename)
        await cb.message.answer(f"Введіть новий номер замість «{apt['number']}» (наприклад, 34/1):")
        await cb.answer()

    @dp.message(PortalForm.apartment_rename)
    async def rename_apartment_save(m: types.Message, state: FSMContext) -> None:
        data = await state.get_data(); apt_id = int(data.get("apartment_id") or 0); number = (m.text or "").strip()
        apt = db.fetch_one("SELECT a.organization_id,o.name org,a.number old FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?", (apt_id,))
        if not apt or not _is_admin(m.from_user.id, apt["org"]):
            await state.clear(); return await m.answer("Немає доступу або квартиру не знайдено.")
        if not re.fullmatch(r"[\wА-Яа-яІіЇїЄєҐґ/-]{1,20}", number):
            return await m.answer("Номер квартири може містити літери, цифри, дефіс і «/», до 20 символів.")
        try:
            db.execute("UPDATE apartments SET number=? WHERE id=?", (number, apt_id))
        except Exception:
            return await m.answer("Такий номер уже є в цьому будинку.")
        _audit(m.from_user.id, "rename_apartment", "apartment", apt_id, apt["org"], f"{apt['old']} -> {number}")
        await state.clear(); await m.answer(f"Номер квартири змінено на {number}. Історія й ключ збережені.")

    @dp.callback_query(F.data.startswith("p:aptinfo:"))
    async def apartment_info_start(cb: types.CallbackQuery, state: FSMContext) -> None:
        apt_id = int(cb.data.rsplit(":", 1)[1])
        apt = db.fetch_one("SELECT a.number,o.name org FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?", (apt_id,))
        if not apt or not _is_admin(cb.from_user.id, apt["org"]):
            return await cb.answer("Немає доступу.", show_alert=True)
        await state.update_data(apartment_id=apt_id); await state.set_state(PortalForm.apartment_info)
        await cb.message.answer(f"Введіть або надішліть текст для кабінету квартири {apt['number']} (надсилайте «-», щоб очистити):")
        await cb.answer()

    @dp.message(PortalForm.apartment_info)
    async def apartment_info_save(m: types.Message, state: FSMContext) -> None:
        data = await state.get_data(); apt_id=int(data.get("apartment_id") or 0)
        apt=db.fetch_one("SELECT o.name org FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?",(apt_id,))
        if not apt or not _is_admin(m.from_user.id,apt["org"]):
            await state.clear(); return await m.answer("Немає доступу.")
        text="" if (m.text or "").strip()=="-" else (m.text or "").strip()
        db.execute("UPDATE apartments SET resident_info=? WHERE id=?",(text,apt_id)); _audit(m.from_user.id,"edit_apartment_info","apartment",apt_id,apt["org"])
        await state.clear(); await m.answer("Інформацію збережено.")

    @dp.callback_query(F.data.startswith("p:aptdoc:"))
    async def apartment_doc_start(cb: types.CallbackQuery,state: FSMContext) -> None:
        apt_id=int(cb.data.rsplit(":",1)[1]); apt=db.fetch_one("SELECT a.number,o.name org FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?",(apt_id,))
        if not apt or not _is_admin(cb.from_user.id,apt["org"]): return await cb.answer("Немає доступу.",show_alert=True)
        await state.update_data(apartment_id=apt_id); await state.set_state(PortalForm.apartment_doc_name)
        await cb.message.answer(f"Введіть назву документа для квартири {apt['number']}:"); await cb.answer()

    @dp.message(PortalForm.apartment_doc_name)
    async def apartment_doc_name_save(m: types.Message,state: FSMContext) -> None:
        name=(m.text or "").strip()
        if not name or len(name)>150: return await m.answer("Введіть назву до 150 символів.")
        await state.update_data(document_name=name); await state.set_state(PortalForm.apartment_doc_file)
        await m.answer("Надішліть PDF або фото. Документ побачить лише ця квартира.")

    @dp.message(PortalForm.apartment_doc_file)
    async def apartment_doc_file_save(m: types.Message,state: FSMContext) -> None:
        data=await state.get_data(); apt_id=int(data.get("apartment_id") or 0)
        apt=db.fetch_one("SELECT a.organization_id,o.name org FROM apartments a JOIN organizations o ON o.id=a.organization_id WHERE a.id=?",(apt_id,))
        if not apt or not _is_admin(m.from_user.id,apt["org"]): await state.clear(); return await m.answer("Немає доступу.")
        if m.photo: file_id=m.photo[-1].file_id
        elif m.document and (m.document.mime_type=="application/pdf" or (m.document.file_name or "").lower().endswith(".pdf")): file_id=m.document.file_id
        else: return await m.answer("Надішліть фото або PDF.")
        doc_id=db.execute("INSERT INTO portal_documents(organization_id,apartment_id,name,file_id,published,for_user,created_by) VALUES(?,?,?,?,1,1,?)",(apt["organization_id"],apt_id,data["document_name"],file_id,m.from_user.id))
        _audit(m.from_user.id,"add_apartment_document","document",doc_id,apt["org"])
        await state.clear(); await m.answer("Документ додано до кабінету цієї квартири.")

    @dp.callback_query(F.data.startswith("p:codes:"))
    async def show_codes(cb: types.CallbackQuery) -> None:
        org_id=int(cb.data.rsplit(":",1)[1]); org=_org(org_id)
        if not org or not _is_admin(cb.from_user.id,org["name"]): return await cb.answer("Немає доступу.",show_alert=True)
        kb=InlineKeyboardMarkup(inline_keyboard=[
            [InlineKeyboardButton(text="🔄 Змінити код мешканців",callback_data=f"p:rotate:{org_id}:user")],
            [InlineKeyboardButton(text="🔄 Змінити код Leadership",callback_data=f"p:rotate:{org_id}:leadership")],
        ])
        await cb.message.answer(f"Коди {org['name']}\nМешканці: {org['resident_code']}\nLeadership: {org['leadership_code']}\n\nЗаміна коду не відключає вже підключених людей.",reply_markup=kb)
        await cb.answer()

    @dp.callback_query(F.data.startswith("p:pinexport:"))
    async def export_apartment_keys(cb: types.CallbackQuery) -> None:
        org_id=int(cb.data.rsplit(":",1)[1]); org=_org(org_id)
        if not org or not _can_export_keys(cb.from_user.id,org["name"]): return await cb.answer("Немає доступу.",show_alert=True)
        payload=_build_pin_xlsx(org_id)
        await cb.message.answer_document(BufferedInputFile(payload,filename=f"Ключі_квартир_{org['name']}.xlsx"),caption="Файл містить ключі входу до квартир. Зберігайте його лише в захищеному місці.")
        _audit(cb.from_user.id,"export_apartment_keys","organization",org_id,org["name"])
        await cb.answer("Файл підготовлено.")

    @dp.callback_query(F.data.startswith("p:rotate:"))
    async def rotate_code(cb: types.CallbackQuery) -> None:
        _,_,org_raw,kind=cb.data.split(":",3); org_id=int(org_raw); org=_org(org_id)
        if not org or not _is_admin(cb.from_user.id,org["name"]): return await cb.answer("Немає доступу.",show_alert=True)
        code=db.create_join_code(); field="resident_code" if kind=="user" else "leadership_code"
        db.execute(f"UPDATE organizations SET {field}=? WHERE id=?",(code,org_id))
        _audit(cb.from_user.id,"rotate_join_code","organization",org_id,org["name"],kind)
        await cb.message.answer(f"Новий код {'мешканців' if kind=='user' else 'Leadership'} для {org['name']}:\n<code>{code}</code>",parse_mode="HTML")
        await cb.answer("Код змінено.")
